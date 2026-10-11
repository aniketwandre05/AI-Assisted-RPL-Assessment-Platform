import os
from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
from backend.database import db
from backend.models.qualification import Qualification, PassingRule

class NQRMasterImporter:
    """
    Imports the official National Qualification Register (NQR) dataset
    from Data_Qualification_Register_11-Oct-2026.xlsx.
    Stores all 1,912 official NSQF qualifications for matching.
    """

    def __init__(self, excel_path: Path):
        self.excel_path = excel_path

    def import_master_register(self, limit: Optional[int] = None) -> Dict[str, Any]:
        if not self.excel_path.exists():
            raise FileNotFoundError(f"NQR register not found at: {self.excel_path}")

        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        sheet = wb.active
        
        headers = [c for c in next(sheet.iter_rows(max_row=1, values_only=True))]
        header_map = {h: i for i, h in enumerate(headers) if h}

        imported_count = 0
        updated_count = 0
        skipped_count = 0
        records = []

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if limit and imported_count >= limit:
                break
                
            q_name = row[header_map.get("Qualification Name", 1)]
            if not q_name:
                continue

            nqr_code = row[header_map.get("NQR Code", 2)]
            sector = row[header_map.get("Sector", 3)] or "General"
            sub_sector = row[header_map.get("Sub-Sector", 4)]
            level_raw = row[header_map.get("NSQF Level", 5)]
            
            try:
                nsqf_level = float(level_raw) if level_raw else 4.0
            except (ValueError, TypeError):
                nsqf_level = 4.0

            min_hours = row[header_map.get("Minimum Notional Hours", 10)]
            max_hours = row[header_map.get("Maximum Notional Hours", 11)]
            approval_date_raw = row[header_map.get("Approval Date", 12)]
            nsqc_no_raw = row[header_map.get("NSQC No.", 13)]
            validity_date_raw = row[header_map.get("Validity Date", 14)]
            awarding_body = row[header_map.get("Name of Awarding Body", 15)]
            th_hours = row[header_map.get("Theory Hours", 18)]
            pr_hours = row[header_map.get("Practical Hours", 19)]
            ojt_hours = row[header_map.get("OJT Hours", 20)]
            emp_hours = row[header_map.get("Employability Skills Hours", 21)]

            # Check if already exists by NQR code or exact title
            existing = None
            if nqr_code:
                existing = Qualification.query.filter_by(nqr_code=str(nqr_code).strip()).first()
            if not existing:
                existing = Qualification.query.filter_by(title=str(q_name).strip()).first()

            if existing:
                # Update metadata if not already filled from detailed PDF
                if not existing.nqr_code and nqr_code:
                    existing.nqr_code = str(nqr_code).strip()
                if not existing.awarding_body and awarding_body:
                    existing.awarding_body = str(awarding_body).strip()
                if not existing.validity_date and validity_date_raw:
                    existing.validity_date = str(validity_date_raw)[:10]
                updated_count += 1
            else:
                qual = Qualification(
                    title=str(q_name).strip(),
                    trade=str(q_name).strip(),
                    sector=str(sector).strip(),
                    sub_sector=str(sub_sector).strip() if sub_sector else None,
                    nsqf_level=nsqf_level,
                    nqr_code=str(nqr_code).strip() if nqr_code else None,
                    awarding_body=str(awarding_body).strip() if awarding_body else None,
                    approval_date=str(approval_date_raw)[:10] if approval_date_raw else None,
                    validity_date=str(validity_date_raw)[:10] if validity_date_raw else None,
                    nsqc_number=str(nsqc_no_raw) if nsqc_no_raw else None,
                    notional_hours_total=float(max_hours) if max_hours else (float(min_hours) if min_hours else None),
                    theory_hours=float(th_hours) if th_hours else None,
                    practical_hours=float(pr_hours) if pr_hours else None,
                    ojt_hours=float(ojt_hours) if ojt_hours else None,
                    employability_hours=float(emp_hours) if emp_hours else None,
                    source_type="NQR_MASTER_REGISTER",
                    is_synthetic=False,
                    is_verified=True,
                    description=f"Official qualification registered with NCVET/NQR under {sector} sector."
                )
                db.session.add(qual)
                imported_count += 1

        db.session.commit()
        return {
            "imported_count": imported_count,
            "updated_count": updated_count,
            "total_processed": imported_count + updated_count + skipped_count,
        }
